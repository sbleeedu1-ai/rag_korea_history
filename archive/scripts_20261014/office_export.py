"""사람이 보기 좋은 제출용 파일로 바꾸기: CSV 여러 개 → 엑셀 시트, 마크다운 → 워드(작업기록 [219]).

CSV·마크다운 원본은 코드와 AI가 읽기 좋은 형식이라 그대로 두고, 산출물 사본만 엑셀·워드로 만든다
(사용자 2026-10-02: "csv는 너가 읽기 좋은 거지 사람은 엑셀이 나아", "정제 규칙서도 워드 문서로").
collect_deliverables.py가 부른다.
"""
import csv
import re
from pathlib import Path

FONT = "맑은 고딕"   # 한글 문서라 본문 글꼴은 맑은 고딕


# ── CSV → 엑셀 ──
def csv_sheets_to_xlsx(sheets: list[tuple[str, "Path | list[list]", str]], out: Path) -> None:
    """sheets: (시트 이름, CSV 경로 또는 줄 목록(첫 줄이 제목 줄), 시트 맨 위 설명 한 줄).
    줄 목록에는 '='로 시작하는 엑셀 수식을 넣을 수 있다(엑셀이 열 때 계산)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    thin = Side(style="thin", color="B7BDC4")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    wb = Workbook()
    wb.remove(wb.active)
    for name, path, note in sheets:
        ws = wb.create_sheet(name[:31])
        rows = path if isinstance(path, list) else list(csv.reader(open(path, encoding="utf-8-sig")))
        ws.append([note])
        ws["A1"].font = Font(name=FONT, italic=True, size=9, color="555555")
        ws.append([])
        for row in rows:
            ws.append([_num(v) if isinstance(v, str) else v for v in row])
        head = 3
        widths = [0] * max(len(r) for r in rows)
        table_end = head + next((i for i, r in enumerate(rows) if not r), len(rows)) - 1   # 빈 줄 뒤(표 아래 설명)는 너비 계산에서 뺀다
        for r in ws.iter_rows(min_row=head, max_row=ws.max_row):
            for c in r:
                c.font = Font(name=FONT, size=9, bold=c.row == head)
                c.border = box
                c.alignment = Alignment(vertical="top", wrap_text=True)
                if c.row == head:
                    c.fill = PatternFill("solid", start_color="E8ECEF")
                if c.value is not None and c.row <= table_end:
                    widths[c.column - 1] = max(widths[c.column - 1], display_width(c.value))
        for i, w in enumerate(widths, 1):   # 글자 길이에 맞추되 긴 글은 50칸에서 줄바꿈
            ws.column_dimensions[get_column_letter(i)].width = min(max(w + 2, 6), 50)
        ws.freeze_panes = ws.cell(row=head + 1, column=1)
        # 표 아래 설명 줄(빈 줄 뒤, 칸 2개): 둘째 칸을 표 너비만큼 병합해 한 줄로 읽히게 한다(사용자 2026-10-02).
        # 엑셀은 병합 칸의 줄 높이를 맞추지 못하므로 글자 폭으로 줄 수를 세어 높이를 정한다.
        ncol = len(rows[0])
        if table_end < ws.max_row and ncol > 2:
            span = sum(ws.column_dimensions[get_column_letter(i)].width or 8 for i in range(2, ncol + 1))
            for r in range(table_end + 2, ws.max_row + 1):
                cell = ws.cell(row=r, column=2)
                if cell.value is None:
                    continue
                ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=ncol)
                lines = max(1, -(-display_width(cell.value) // int(span)))
                ws.row_dimensions[r].height = 15 * lines
                if r == table_end + 2:   # 설명 표의 제목 줄
                    for c in (ws.cell(row=r, column=1), cell):
                        c.font = Font(name=FONT, size=9, bold=True)
                        c.fill = PatternFill("solid", start_color="E8ECEF")
    wb.save(out)


def display_width(value) -> int:
    """칸에 보이는 글자 폭. 한글·한자는 2칸, 나머지는 1칸. 수식은 결과(짧은 숫자·비율)만 보이므로 12칸으로 친다.
    여러 줄 글은 가장 긴 줄 기준(사용자 2026-10-02: 수식 글자 길이 때문에 칸이 너무 넓었음)."""
    s = str(value)
    if s.startswith("="):
        return 12
    return max((sum(2 if ord(ch) > 0x2E7F else 1 for ch in line) for line in s.splitlines() or [""]), default=0)


def _num(v: str):
    """숫자 칸은 숫자로(엑셀에서 정렬·합계가 되게). 앞자리 0이 뜻을 갖는 값은 글자로 둔다."""
    if re.fullmatch(r"-?(0|[1-9]\d*)(\.\d+)?", v or ""):
        return float(v) if "." in v else int(v)
    return v


# ── 마크다운 → 워드 ──
_INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def _runs(par, text: str, bold: bool = False):
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            _runs(par, part[2:-2], bold=True)
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.bold = bold
        elif part.startswith("["):
            r = par.add_run(re.match(r"\[([^\]]+)\]", part).group(1))   # 저장소 안 링크는 글자만 남김
            r.bold = bold
        else:
            r = par.add_run(part)
            r.bold = bold


def _cells(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def md_to_docx(src: Path, out: Path, note: str) -> dict:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    doc = Document()
    for style in ("Normal", "Heading 1", "Heading 2", "Heading 3", "Heading 4", "List Bullet", "List Number"):
        st = doc.styles[style]
        st.font.name = FONT
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    doc.styles["Normal"].font.size = Pt(10)
    p = doc.add_paragraph()
    r = p.add_run(note)
    r.italic, r.font.size, r.font.color.rgb = True, Pt(8.5), RGBColor(0x55, 0x55, 0x55)

    lines = src.read_text(encoding="utf-8").splitlines()
    i, para, stats = 0, [], {"제목": 0, "표": 0, "목록": 0, "코드": 0}

    def flush():
        if para:
            _runs(doc.add_paragraph(), " ".join(para))
            para.clear()

    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if s.startswith("```"):
            flush()
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                cp = doc.add_paragraph()
                cr = cp.add_run(lines[i])
                cr.font.name, cr.font.size = "Consolas", Pt(8.5)
                cp.paragraph_format.space_after = Pt(0)
                i += 1
            stats["코드"] += 1
        elif s.startswith("#"):
            flush()
            level = len(s) - len(s.lstrip("#"))
            doc.add_heading(s.lstrip("#").strip(), level=min(level, 4))
            stats["제목"] += 1
        elif s.startswith("|"):
            flush()
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not re.fullmatch(r"\|?[\s:\-|]+\|?", lines[i].strip()):
                    block.append(_cells(lines[i]))
                i += 1
            ncol = max(len(r) for r in block)
            table = doc.add_table(rows=len(block), cols=ncol)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.LEFT
            for ri, row in enumerate(block):
                for ci in range(ncol):
                    cell = table.cell(ri, ci)
                    cell.text = ""
                    _runs(cell.paragraphs[0], row[ci] if ci < len(row) else "", bold=ri == 0)
                    for run in cell.paragraphs[0].runs:
                        run.font.size = Pt(8.5)
                    if ri == 0:
                        shd = OxmlElement("w:shd")
                        shd.set(qn("w:val"), "clear")
                        shd.set(qn("w:fill"), "E8ECEF")
                        cell._tc.get_or_add_tcPr().append(shd)
            doc.add_paragraph()
            stats["표"] += 1
            continue
        elif re.match(r"^\s*[-*] ", line):
            flush()
            _runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*] ", "", line))
            stats["목록"] += 1
        elif re.match(r"^\s*\d+\. ", line):
            flush()
            _runs(doc.add_paragraph(style="List Number"), re.sub(r"^\s*\d+\. ", "", line))
            stats["목록"] += 1
        elif s in ("---", "***") or not s:
            flush()
        else:
            para.append(s)
        i += 1
    flush()
    doc.save(out)
    return stats
